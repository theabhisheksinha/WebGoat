from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt
from docx.oxml.ns import qn


OUTPUT_PATH = r"c:\Apps\WebGoat_v3\docs\WebGoat_v3_Technical_Assessment_Report.docx"


def set_default_font(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    style.font.size = Pt(10)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_paragraph(doc: Document, text: str = ""):
    return doc.add_paragraph(text)


def add_bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def add_code_block(doc: Document, title: str, path: str, lines: str, snippet: str) -> None:
    p = doc.add_paragraph()
    p.add_run(f"{title}\n").bold = True
    code = p.add_run(f"{path}:{lines}\n{snippet}")
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code.font.size = Pt(9)


def add_table(doc: Document, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    return table


doc = Document()
set_default_font(doc)

add_heading(doc, "Technical Assessment Report: WebGoat_v3", 0)
add_paragraph(doc, "Application: WebGoat_v3")
add_paragraph(doc, "Assessment basis: CAST Imaging MCP evidence and verified local source inspection.")

# 1. Executive Summary
add_heading(doc, "1. Executive Summary", 1)
add_paragraph(
    doc,
    "WebGoat_v3 is a legacy Java servlet/JSP monolith with significant modernization, security, "
    "and database portability concerns. CAST Imaging indicates high maintainability debt, notable "
    "security and reliability exposure, and low cloud readiness. The strongest migration path is "
    "replatform plus selective refactoring toward AWS PaaS with strict layering, externalized state, "
    "and PostgreSQL on Amazon RDS."
)
add_bullet(doc, "Overall risk profile: High")
add_bullet(doc, "Cloud/PaaS maturity: Low")
add_bullet(doc, "Modernization strategy: Replatform first, then refactor hot spots")
add_bullet(doc, "Primary blockers: embedded HSQLDB, runtime schema creation, servlet session state, Axis/JAX-RPC/CORBA legacy stack")

# 2. System Topology
add_heading(doc, "2. System Topology", 1)
add_paragraph(
    doc,
    "CAST Imaging component analysis shows the application organized around Web Interaction, Logic Services, "
    "and Database Services. Evidence from the component links indicates reverse coupling from Logic Services "
    "back into Web Interaction, which violates the target unidirectional layering expected for a compliant architecture."
)
add_table(
    doc,
    ["Evidence Type", "Finding"],
    [
        ["Component graph", "Web Interaction -> Logic Services"],
        ["Component graph", "Logic Services -> Database Services"],
        ["Component graph", "Logic Services -> Web Interaction"],
        ["API inventory", "6 public servlet endpoints in Imaging"],
        ["Transaction profile", "Largest servlet transactions reach size 2233 objects"],
    ],
)
add_code_block(
    doc,
    "Public servlet entry point",
    r"src-5.2\webgoat\JavaSource\org\owasp\webgoat\HammerHead.java",
    "10-24",
    "import javax.servlet.ServletContext;\n"
    "import javax.servlet.ServletException;\n"
    "import javax.servlet.http.HttpServlet;\n"
    "...\n"
    "import org.owasp.webgoat.session.WebSession;\n"
    "import org.owasp.webgoat.session.WebgoatContext;",
)
add_code_block(
    doc,
    "Controller servlet",
    r"src-5.2\webgoat\JavaSource\org\owasp\webgoat\servlets\Controller.java",
    "1-20",
    "package org.owasp.webgoat.servlets;\n"
    "import javax.servlet.http.HttpServlet;\n"
    "...\n"
    "public class Controller extends HttpServlet",
)

# 3. Technical Debt (ISO-5055)
add_heading(doc, "3. Technical Debt (ISO-5055)", 1)
add_paragraph(
    doc,
    "CAST Imaging reports 1,544 ISO-5055-impacted objects across four characteristics."
)
add_table(
    doc,
    ["Characteristic", "Impacted Objects"],
    [
        ["Maintainability", "743"],
        ["Efficiency", "284"],
        ["Reliability", "233"],
        ["Security", "284"],
    ],
)
add_table(
    doc,
    ["Characteristic", "Representative Weakness", "Count"],
    [
        ["Maintainability", "CWE-1048 Invokable Control Element with Large Number of Outward Calls", "441"],
        ["Maintainability", "CWE-1041 Use of Redundant Code", "185"],
        ["Reliability", "CWE-391 Unchecked Error Condition", "48"],
        ["Reliability", "CWE-252 Unchecked Return Value", "38"],
        ["Efficiency", "CWE-1046 String Concatenation for Immutable Text", "94"],
        ["Efficiency", "CWE-1049 Excessive Data Query Operations", "58"],
        ["Security", "CWE-404 Improper Resource Shutdown or Release", "85"],
        ["Security", "CWE-477 Use of Obsolete Function", "50"],
    ],
)
add_paragraph(doc, "Critical violation locations:")
add_paragraph(doc, "**Data not found from skills and CAST mcp**")
add_paragraph(
    doc,
    "Reason: CAST returned weakness counts for ISO-5055 families, but occurrence-level location payloads were not returned "
    "for representative weakness IDs such as CWE-1048, CWE-404, and CWE-259 in this application."
)

# 4. Modernization & Cloud Score
add_heading(doc, "4. Modernization & Cloud Score", 1)
add_paragraph(
    doc,
    "CAST Imaging reported 22 cloud detection patterns with 2 Critical and 5 High findings. "
    "The application is not ready for resilient PaaS deployment without structural remediation."
)
add_table(
    doc,
    ["Finding", "Severity", "Impacted Objects"],
    [
        ["Use of an unsecured data string", "Critical", "72"],
        ["Deprecated language or framework versions", "Critical", "1"],
        ["Using stateful session", "High", "4"],
        ["Using CORBA", "High", "1"],
        ["Using JAX-RPC technology", "High", "1"],
        ["Using Java RMI", "High", "1"],
        ["Using file system", "Low", "18"],
        ["Avoid hardcoded HTTP URLs", "Low", "76"],
    ],
)
add_code_block(
    doc,
    "Embedded HSQLDB configuration",
    r"src-5.2\webgoat\WEB-INF\web.xml",
    "160-174",
    "<param-name>DatabaseDriver</param-name>\n"
    "<param-value>\n"
    "    org.hsqldb.jdbcDriver\n"
    "</param-value>\n"
    "...\n"
    "<param-value>\n"
    "    jdbc:hsqldb:mem:${USER}\n"
    "</param-value>",
)
add_code_block(
    doc,
    "Legacy remoting stack",
    r"src-5.2\webgoat\JavaSource\org\owasp\webgoat\lessons\WSDLScanning.java",
    "8-20",
    "import java.rmi.RemoteException;\n"
    "import javax.xml.rpc.ParameterMode;\n"
    "import javax.xml.rpc.ServiceException;\n"
    "import org.apache.axis.client.Call;\n"
    "import org.apache.axis.client.Service;\n"
    "import org.apache.axis.encoding.XMLType;",
)
add_code_block(
    doc,
    "Legacy CORBA reference",
    r"src-5.2\webgoat\JavaSource\org\owasp\webgoat\lessons\CrossSiteScripting\ViewProfile.java",
    "18-20",
    "import org.owasp.webgoat.util.HtmlEncoder;\n"
    "\n"
    "import com.sun.corba.se.spi.activation.Server;",
)

# 5. Database Strategy
add_heading(doc, "5. Database Strategy", 1)
add_paragraph(
    doc,
    "The current runtime uses embedded HSQLDB while the repository contains Oracle, SQL Server, and PostgreSQL schema artifacts. "
    "The recommended target is Amazon RDS for PostgreSQL with schema migrations managed outside the application runtime."
)
add_bullet(doc, "Current runtime DB: HSQLDB in-memory per user")
add_bullet(doc, "Target DB: Amazon RDS PostgreSQL")
add_bullet(doc, "Migration pattern: baseline PostgreSQL schema + Flyway/Liquibase + controlled seed data")
add_bullet(doc, "Do not retain runtime schema creation in the application")
add_code_block(
    doc,
    "Direct JDBC and vendor-specific connection logic",
    r"src-5.2\webgoat\JavaSource\org\owasp\webgoat\session\DatabaseUtilities.java",
    "90-107",
    "private static Connection makeConnection(String user, WebgoatContext context) throws ClassNotFoundException,\n"
    "        SQLException\n"
    "{\n"
    "    Class.forName(context.getDatabaseDriver());\n"
    "    if (context.getDatabaseConnectionString().contains(\"hsqldb\")) return getHsqldbConnection(user, context);\n"
    "    ...\n"
    "    return DriverManager.getConnection(url, userPrefix + \"_\" + user, password);\n"
    "}",
)
add_code_block(
    doc,
    "Hardcoded seed credentials in runtime DB builder",
    r"src-5.2\webgoat\JavaSource\org\owasp\webgoat\session\CreateDB.java",
    "163-167",
    "String insertData1 = \"INSERT INTO user_system_data VALUES ('101','jsnow','passwd1', '')\";\n"
    "String insertData2 = \"INSERT INTO user_system_data VALUES ('102','jdoe','passwd2', '')\";\n"
    "String insertData3 = \"INSERT INTO user_system_data VALUES ('103','jplane','passwd3', '')\";\n"
    "String insertData4 = \"INSERT INTO user_system_data VALUES ('104','jeff','jeff', '')\";\n"
    "String insertData5 = \"INSERT INTO user_system_data VALUES ('105','dave','dave', '')\";",
)

# 6. Quality & Governance
add_heading(doc, "6. Quality & Governance", 1)
add_paragraph(
    doc,
    "The target architecture must enforce strict layering. CAST Imaging component links already show a reverse dependency from Logic Services "
    "to Web Interaction, which is a governance defect relative to the desired architecture."
)
add_table(
    doc,
    ["Governance Rule", "Current State", "Assessment"],
    [
        ["Web layer may call services", "Observed", "Compliant"],
        ["Services may call database services", "Observed", "Compliant"],
        ["Services must not depend on web layer", "Observed reverse link", "Violation"],
        ["No local filesystem for durable state", "File-based lesson tracking present", "Violation"],
        ["No runtime schema creation", "CreateDB invoked at runtime", "Violation"],
    ],
)
add_code_block(
    doc,
    "Servlet-based stateful boundary",
    r"src-5.2\webgoat\JavaSource\org\owasp\webgoat\HammerHead.java",
    "12-24",
    "import javax.servlet.http.HttpServlet;\n"
    "import javax.servlet.http.HttpServletRequest;\n"
    "import javax.servlet.http.HttpServletResponse;\n"
    "import javax.servlet.http.HttpSession;\n"
    "...\n"
    "import org.owasp.webgoat.session.WebSession;",
)
add_paragraph(
    doc,
    "External dependency compliance:"
)
add_paragraph(doc, "**Data not found from skills and CAST mcp**")

# 7. Security Posture
add_heading(doc, "7. Security Posture", 1)
add_paragraph(
    doc,
    "CAST Imaging returned 29 CVE findings in third-party packages and several structural security weakness families. "
    "The most material verified risk is the continued use of Apache Axis 1.x and adjacent legacy libraries."
)
add_table(
    doc,
    ["CVE", "Severity", "Summary"],
    [
        ["CVE-2023-40743", "Critical", "Axis 1.x unsafe ServiceFactory lookup can expose SSRF/DoS/RCE patterns"],
        ["CVE-2023-51441", "High", "Axis admin service input validation issue enabling SSRF"],
        ["CVE-2019-0227", "High", "Axis 1.4 SSRF vulnerability"],
        ["CVE-2018-8032", "Medium", "Axis 1.x XSS in default servlet/services"],
        ["CVE-2025-48734", "High", "Commons BeanUtils vulnerability present in dependency set"],
    ],
)
add_code_block(
    doc,
    "Hard-coded authentication example",
    r"src-5.2\webgoat\JavaSource\org\owasp\webgoat\lessons\WeakAuthenticationCookie.java",
    "126-138",
    "String password = s.getParser().getStringParameter(PASSWORD, \"\");\n"
    "...\n"
    "if (username.equals(\"webgoat\") && password.equals(\"webgoat\"))\n"
    "{\n"
    "    loginID = encode(\"webgoat12345\");\n"
    "}\n"
    "else if (username.equals(\"aspect\") && password.equals(\"aspect\"))",
)
add_code_block(
    doc,
    "Legacy SOAP stack driving CVE exposure",
    r"src-5.2\webgoat\JavaSource\org\owasp\webgoat\lessons\WSDLScanning.java",
    "16-20",
    "import javax.xml.rpc.ParameterMode;\n"
    "import javax.xml.rpc.ServiceException;\n"
    "import org.apache.axis.client.Call;\n"
    "import org.apache.axis.client.Service;\n"
    "import org.apache.axis.encoding.XMLType;",
)
add_paragraph(doc, "Source and sink line numbers for data-flow risks:")
add_paragraph(doc, "**Data not found from skills and CAST mcp**")

# 8. Modernization Recommendation
add_heading(doc, "8. Modernization Recommendation", 1)
add_paragraph(
    doc,
    "Derived from the verified evidence in Sections 2-7, the recommended path is a phased AWS modernization "
    "toward Elastic Beanstalk, Amazon RDS PostgreSQL, ElastiCache Redis, Amazon S3, and centralized secrets/configuration."
)
add_bullet(doc, "Replace HSQLDB and runtime CreateDB logic with PostgreSQL migrations managed by Flyway or Liquibase")
add_bullet(doc, "Refactor servlet/JSP flows so business logic is moved behind application services")
add_bullet(doc, "Remove Axis 1.x, JAX-RPC, CORBA, and RMI dependencies from the runtime path")
add_bullet(doc, "Externalize session state and file-backed lesson state")
add_bullet(doc, "Enforce unidirectional dependencies: Web -> Service -> Data Access -> Database")

# 9. Disclaimer
add_heading(doc, "9. Disclaimer", 1)
add_paragraph(
    doc,
    "This document is AI-generated based on structural analysis and requires human review and validation."
)

doc.save(OUTPUT_PATH)
print(OUTPUT_PATH)
